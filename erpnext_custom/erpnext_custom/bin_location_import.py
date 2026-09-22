"""Impor layout rak/bin legacy (Bin Location.xlsx) ke master Rack + Bin Location.

Jalankan dari host:

    python erpnext_custom/erpnext_custom/bin_location_import.py "D:/System_ERPNext/Bin Location.xlsx" > bins.json
    docker cp bins.json erp_cakra-backend-1:/tmp/bins.json
    docker exec -i erp_cakra-backend-1 bench --site erp.localhost console
    >>> from erpnext_custom.bin_location_import import run; run("/tmp/bins.json")

Kolom Used SENGAJA tidak diimpor: isi bin dihitung dari Item Bin Qty, jadi tidak
pernah basi. Kapasitas raksasa (>= 1e8) dianggap "tanpa batas" -> 0.
"""

import json
import sys

UNLIMITED = 1e8


def _s(value):
	"""Sel Excel jadi teks. Sebagian kode bin ('0101') tersimpan sebagai ANGKA di
	file aslinya -- kode penuh yang benar ada di kolom Generate Code, jadi yang di
	sini cukup dijadikan teks apa adanya."""
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return str(value).strip()


def read_xlsx(path):
	"""Baca sheet Data jadi list dict. Hanya butuh openpyxl di HOST, bukan di container."""
	import openpyxl

	wb = openpyxl.load_workbook(path, read_only=True)
	rows = list(wb["Data"].iter_rows(min_row=4, values_only=True))
	out = []
	for r in rows:
		if not r or not r[1]:
			continue
		out.append(
			{
				"warehouse_name": _s(r[1]),
				"parent_rack": _s(r[2]) or None,
				"code": _s(r[4]),
				"gen_code": _s(r[5]) or None,
				"level": _s(r[6]) or None,
				"uom": _s(r[7]) or None,
				"capacity": float(r[8] or 0),
				"bulk_area": _s(r[12]).lower() == "yes",
				"customer_area": _s(r[13]).lower() == "yes",
				"disabled": _s(r[14]).lower() != "active",
			}
		)
	return out


# Peruntukan bin di file Rev3 ditulis dengan ejaan gudang; ini padanannya di master UOM.
# "Blank" berarti bin itu tidak dipesan untuk kemasan tertentu, bukan nama satuan.
_UOM_REV3 = {"PCS": "Pcs", "ZAK": "Zak", "Drum": "Drum", "Box": "Box", "Roll": "Roll",
             "Batang": "Batang", "Blank": None}


def read_map_rev3(path, sheet="Layout 5L Map Rev3", warehouse="Gudang Jakarta"):
	"""Baca tabel Master BIN di file Cakraindo Material Mapping Check.xlsx.

	Satu LORONG di file itu sebenarnya DUA baris rak yang berhadapan: bay ganjil di
	sisi satu, genap di sisi lain (kolom Category). Jadi rak yang dibuat = "AA Ganjil"
	dan "AA Genap", supaya denah 2D-nya sama dengan gudang aslinya dan tiap kotak
	benar-benar satu baris rak.

	Kolom "Capacity Bin" TIDAK diimpor sebagai kapasitas berat: isinya 1, artinya satu
	palet per bin, bukan 1 kg. Kapasitas kg/m3 belum ada di file ini.
	"""
	import openpyxl

	wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
	racks, bins = {}, []
	for r in wb[sheet].iter_rows(min_row=3, min_col=2, max_col=9, values_only=True):
		if not r[0]:
			continue
		bin_code, aisle, _bay, level, kategori, _cap, status, uom = (_s(x) for x in r)
		rak = "{0} {1}".format(aisle, kategori)
		racks.setdefault(rak, {"warehouse_name": warehouse, "parent_rack": None, "code": rak,
		                       "gen_code": None, "level": None, "uom": None, "capacity": 0,
		                       "bulk_area": False, "customer_area": False, "disabled": False})
		bins.append(
			{
				"warehouse_name": warehouse,
				"parent_rack": rak,
				"code": bin_code,
				"gen_code": None,
				"level": level or None,
				"uom": _UOM_REV3.get(uom, uom or None),
				"capacity": 0,
				"bulk_area": False,
				"customer_area": status == "Storage For Customer",
				"disabled": False,
			}
		)
	return list(racks.values()) + bins


def run(path, abbr=None, dry_run=False):
	"""Buat Rack + Bin Location dari file JSON hasil read_xlsx. Aman dijalankan ulang."""
	import frappe

	data = json.load(open(path, encoding="utf8"))
	if abbr is None:
		abbr = frappe.get_all("Company", limit=1, pluck="abbr")[0]

	def gudang_of(row):
		return "{0} - {1}".format(row["warehouse_name"], abbr)

	missing, made_rack, made_bin, skipped = set(), 0, 0, 0

	# 1. rak (baris tanpa Parent Rack)
	for row in data:
		if row["parent_rack"]:
			continue
		gudang = gudang_of(row)
		if not frappe.db.exists("Warehouse", gudang):
			missing.add(gudang)
			continue
		if frappe.db.exists("Rack", {"gudang": gudang, "rack_code": row["code"]}):
			continue
		if not dry_run:
			frappe.get_doc(
				{
					"doctype": "Rack",
					"gudang": gudang,
					"rack_code": row["code"],
					"disabled": int(row["disabled"]),
				}
			).insert(ignore_permissions=True)
		made_rack += 1

	# 2. bin
	for row in data:
		if not row["parent_rack"]:
			continue
		gudang = gudang_of(row)
		if not frappe.db.exists("Warehouse", gudang):
			missing.add(gudang)
			continue
		rack = frappe.db.exists("Rack", {"gudang": gudang, "rack_code": row["parent_rack"]})
		if not rack:
			# 13 bin di gudang sparepart menunjuk rak "Staging Area" yang tidak punya
			# baris sendiri di file legacy -- raknya dibuatkan, jangan dibuang binnya.
			if dry_run:
				continue
			rack = frappe.get_doc(
				{"doctype": "Rack", "gudang": gudang, "rack_code": row["parent_rack"]}
			).insert(ignore_permissions=True).name
			made_rack += 1
		bin_code = row["gen_code"] or row["code"]
		if frappe.db.exists("Bin Location", {"gudang": gudang, "bin_code": bin_code}):
			skipped += 1
			continue
		uom = row["uom"] if row["uom"] and frappe.db.exists("UOM", row["uom"]) else None
		if not dry_run:
			frappe.get_doc(
				{
					"doctype": "Bin Location",
					"rack": rack,
					"bin_code": bin_code,
					"level": row["level"],
					"default_uom": uom,
					"capacity_weight": 0 if row["capacity"] >= UNLIMITED else row["capacity"],
					"bulk_area": int(row["bulk_area"]),
					"customer_area": int(row["customer_area"]),
					"disabled": int(row["disabled"]),
				}
			).insert(ignore_permissions=True)
		made_bin += 1

	if not dry_run:
		frappe.db.commit()
	return {
		"rack": made_rack,
		"bin": made_bin,
		"sudah_ada": skipped,
		"gudang_tidak_ketemu": sorted(missing),
	}


def seed_rack_levels(gudang=None):
	"""Siapkan baris Tingkat kosong di master Rack, satu per tingkat yang benar-benar
	ada bin-nya. Tinggal diketik kapasitasnya di form Rak.

	Kapasitas volume dan panjang slot tidak ada di file legacy, jadi diisi per
	TINGKAT (16 rak x ~5 tingkat) ketimbang per bin (546). Bin yang kapasitasnya
	kosong otomatis ikut baris tingkat ini -- lihat bin_layout._own_caps.

	Rak yang sudah punya baris Tingkat DILEWATI, jadi aman dijalankan berulang.
	"""
	import frappe

	filters = {"gudang": gudang} if gudang else {}
	dibuat, dilewati = {}, 0
	for rack in frappe.get_all("Rack", filters=filters, fields=["name", "kind"]):
		if rack.kind not in ("Rak", "Staging"):
			continue
		doc = frappe.get_doc("Rack", rack.name)
		if doc.get("levels"):
			dilewati += 1
			continue
		tingkat = sorted(
			{
				(lv or "").strip().upper()
				for lv in frappe.get_all("Bin Location", filters={"rack": rack.name}, pluck="level")
				if (lv or "").strip()
			}
		)
		if not tingkat:
			continue
		for lv in tingkat:
			doc.append("levels", {"level": lv})
		doc.save(ignore_permissions=True)
		dibuat[rack.name] = tingkat
	frappe.db.commit()
	return {"dibuat": dibuat, "sudah_punya": dilewati}


if __name__ == "__main__":
	# argumen kedua "rev3" = file mapping gudang Jakarta yang baru, tanpa itu file legacy
	baca = read_map_rev3 if sys.argv[2:3] == ["rev3"] else read_xlsx
	json.dump(baca(sys.argv[1]), sys.stdout, ensure_ascii=False, indent=1)

"""Buka lampiran .xlsx apa adanya (format ikut file), edit di grid, simpan balik ke file yang sama.

ponytail: openpyxl menulis ulang file -> chart, gambar, pivot & komentar sel bisa hilang,
dan nilai hasil rumus jadi basi sampai file dibuka Excel. Kalau itu jadi masalah,
ganti jalur simpan ke LibreOffice headless (soffice --convert-to xlsx) yang menghitung ulang.
"""

import hashlib
import json
import os
import re
from datetime import date, datetime, time

import frappe
from frappe import _
from frappe.utils import cint

MAX_ROWS = 2000
MAX_COLS = 100
EDITABLE_EXT = (".xlsx", ".xlsm")

_LITERAL = re.compile(r'"([^"]*)"|\[\$([^\]\-]*)[^\]]*\]')
_NUMERIC = re.compile(r"-?\d+(\.\d+)?([eE][+-]?\d+)?")


def _get_file(file_name: str, ptype: str = "read"):
	doc = frappe.get_doc("File", file_name)
	if doc.attached_to_doctype:
		# izin ikut dokumen induknya (Inquiry / Quotation / dst)
		frappe.has_permission(doc.attached_to_doctype, ptype, doc=doc.attached_to_name, throw=True)
	elif ptype != "read" and doc.owner != frappe.session.user:
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	if not (doc.file_name or "").lower().endswith(EDITABLE_EXT):
		frappe.throw(_("Hanya file .xlsx / .xlsm yang bisa dibuka di editor"))

	return doc


def _color(c):
	rgb = getattr(c, "rgb", None)
	if isinstance(rgb, str) and len(rgb) == 8 and rgb[:2] != "00":
		return "#" + rgb[2:]
	return None


def _plain(v):
	if v is None:
		return ""
	if isinstance(v, bool):
		return "TRUE" if v else "FALSE"
	if isinstance(v, datetime):
		return v.strftime("%d/%m/%Y") if v.time() == time(0, 0) else v.strftime("%d/%m/%Y %H:%M")
	if isinstance(v, date):
		return v.strftime("%d/%m/%Y")
	if isinstance(v, time):
		return v.strftime("%H:%M")
	if isinstance(v, float) and v.is_integer():
		return str(int(v))
	return str(v)


def _display(v, fmt):
	"""ponytail: subset format angka Excel (literal, ribuan, desimal, persen), pemisah gaya en-US.
	Format kondisional & tanggal kustom jatuh balik ke _plain."""
	if isinstance(v, bool) or not isinstance(v, (int, float)):
		return _plain(v)

	body = (fmt or "General").split(";")[0]
	if body in ("", "General", "@"):
		return _plain(v)

	try:
		prefix = "".join(a or b for a, b in _LITERAL.findall(body))
		core = _LITERAL.sub("", body)
		if "%" in core:
			v = v * 100
		dec = len(re.sub(r"[^0#]", "", core.split(".")[1])) if "." in core else 0
		text = f"{v:,.{dec}f}" if "," in core else f"{v:.{dec}f}"
		return f"{prefix}{text}{'%' if '%' in core else ''}".strip()
	except Exception:
		return _plain(v)


def _style(cell):
	s = {}
	f = cell.font
	if f.bold:
		s["b"] = 1
	if f.italic:
		s["i"] = 1
	if f.underline:
		s["u"] = 1
	if f.size and float(f.size) != 11:
		s["sz"] = float(f.size)
	if f.name and f.name != "Calibri":
		s["ff"] = f.name

	color = _color(f.color) if f.color else None
	if color and color != "#000000":
		s["c"] = color

	if cell.fill is not None and cell.fill.patternType == "solid":
		bg = _color(cell.fill.fgColor)
		if bg and bg != "#FFFFFF":
			s["bg"] = bg

	a = cell.alignment
	if a.horizontal:
		s["ha"] = a.horizontal
	if a.vertical and a.vertical != "bottom":
		s["va"] = a.vertical
	if a.wrap_text:
		s["w"] = 1

	bd = {}
	for side in ("top", "right", "bottom", "left"):
		edge = getattr(cell.border, side, None)
		if edge and edge.style:
			bd[side] = (_color(edge.color) if edge.color else None) or "#64748b"
	if bd:
		s["bd"] = bd

	return s


def _sheet_payload(wsf, wsv):
	from openpyxl.utils import get_column_letter

	nrows = min(wsf.max_row or 1, MAX_ROWS)
	ncols = min(wsf.max_column or 1, MAX_COLS)

	default_w = wsf.sheet_format.defaultColWidth or 8.43
	default_h = wsf.sheet_format.defaultRowHeight or 15

	widths, hidden_cols = [], []
	for c in range(1, ncols + 1):
		d = wsf.column_dimensions.get(get_column_letter(c))
		if d is not None and d.hidden:
			hidden_cols.append(c)
		widths.append(round((d.width if d is not None and d.width else default_w) * 7 + 5))

	heights = []
	for r in range(1, nrows + 1):
		d = wsf.row_dimensions.get(r)
		heights.append(round((d.height if d is not None and d.height else default_h) * 4 / 3))

	merges, covered = [], set()
	for rng in wsf.merged_cells.ranges:
		if rng.min_row > nrows or rng.min_col > ncols:
			continue
		merges.append(
			{
				"r": rng.min_row,
				"c": rng.min_col,
				"rs": min(rng.max_row, nrows) - rng.min_row + 1,
				"cs": min(rng.max_col, ncols) - rng.min_col + 1,
			}
		)
		for rr in range(rng.min_row, min(rng.max_row, nrows) + 1):
			for cc in range(rng.min_col, min(rng.max_col, ncols) + 1):
				if (rr, cc) != (rng.min_row, rng.min_col):
					covered.add((rr, cc))

	cells = []
	for r in range(1, nrows + 1):
		row = []
		for c in range(1, ncols + 1):
			if (r, c) in covered:
				row.append(0)  # ditelan sel merge di kiri/atasnya
				continue
			fc = wsf.cell(row=r, column=c)
			raw = wsv.cell(row=r, column=c).value
			formula = fc.value if isinstance(fc.value, str) and fc.value.startswith("=") else None
			if raw is None and formula:
				raw = formula  # rumus belum pernah dihitung Excel
			style = _style(fc)
			if isinstance(raw, (int, float, date, time)) and not isinstance(raw, bool):
				style.setdefault("ha", "right")  # default Excel: angka & tanggal rata kanan
			if raw is None and not style:
				row.append(None)
				continue
			cell = {"v": _display(raw, fc.number_format)}
			if formula:
				cell["f"] = formula
			if style:
				cell["s"] = style
			row.append(cell)
		cells.append(row)

	return {
		"cells": cells,
		"merges": merges,
		"widths": widths,
		"heights": heights,
		"hidden_cols": hidden_cols,
		"rows": nrows,
		"cols": ncols,
		"truncated": (wsf.max_row or 0) > MAX_ROWS or (wsf.max_column or 0) > MAX_COLS,
	}


@frappe.whitelist()
def read(file_name: str, sheet: str | None = None):
	from openpyxl import load_workbook

	doc = _get_file(file_name)
	path = doc.get_full_path()

	wbv = load_workbook(path, data_only=True)  # nilai hasil rumus (cache Excel)
	wbf = load_workbook(path, data_only=False)  # rumus + style
	name = sheet if sheet in wbf.sheetnames else wbf.sheetnames[0]

	payload = _sheet_payload(wbf[name], wbv[name])
	payload.update(
		{
			"file": doc.name,
			"file_name": doc.file_name,
			"file_url": doc.file_url,
			"sheet": name,
			"sheets": wbf.sheetnames,
		}
	)
	wbv.close()
	wbf.close()
	return payload


def _coerce(v):
	"""Teks dari grid -> tipe Excel. Angka berawalan 0 (kode, NPWP) tetap teks."""
	if v is None:
		return None
	s = str(v).strip()
	if not s:
		return None
	if s.startswith("="):
		return s
	if _NUMERIC.fullmatch(s) and not (len(s) > 1 and s[0] == "0" and s[1] != "."):
		return float(s) if ("." in s or "e" in s.lower()) else int(s)
	return s


@frappe.whitelist()
def save(file_name: str, sheet: str, changes):
	from openpyxl import load_workbook
	from openpyxl.utils import get_column_letter

	doc = _get_file(file_name, "write")
	if isinstance(changes, str):
		changes = json.loads(changes)
	if not changes:
		return read(file_name, sheet)

	path = doc.get_full_path()
	wb = load_workbook(path, data_only=False)  # rumus & style sel lain tetap utuh
	if sheet not in wb.sheetnames:
		frappe.throw(_("Sheet {0} tidak ada di file ini").format(sheet))
	ws = wb[sheet]

	applied = []
	for ch in changes:
		r, c = cint(ch.get("r")), cint(ch.get("c"))
		if r < 1 or c < 1 or r > MAX_ROWS or c > MAX_COLS:
			continue
		cell = ws.cell(row=r, column=c)
		if cell.__class__.__name__ == "MergedCell":
			continue
		before = cell.value
		after = _coerce(ch.get("v"))
		if before == after:
			continue
		cell.value = after
		applied.append((f"{get_column_letter(c)}{r}", before, after))

	if not applied:
		wb.close()
		return read(file_name, sheet)

	wb.save(path)
	wb.close()

	with open(path, "rb") as f:
		content_hash = hashlib.md5(f.read(), usedforsecurity=False).hexdigest()
	doc.db_set({"file_size": os.stat(path).st_size, "content_hash": content_hash})

	_log_changes(doc, sheet, applied)

	if doc.attached_to_doctype and doc.attached_to_name:
		frappe.get_doc(doc.attached_to_doctype, doc.attached_to_name).add_comment(
			"Comment",
			_("Mengubah {0} sel di {1} ({2})").format(len(applied), doc.file_name, sheet),
		)

	return read(file_name, sheet)


def _log_changes(doc, sheet: str, applied: list):
	"""Satu baris CRM Spreadsheet Log per sel; batch menandai satu kali Save."""
	batch = frappe.generate_hash(length=10)
	for cell, before, after in applied:
		frappe.get_doc(
			{
				"doctype": "CRM Spreadsheet Log",
				"file": doc.name,
				"file_name": doc.file_name,
				"attached_to_doctype": doc.attached_to_doctype,
				"attached_to_name": doc.attached_to_name,
				"sheet": sheet,
				"cell": cell,
				"old_value": _plain(before),
				"new_value": _plain(after),
				"batch": batch,
			}
		).insert(ignore_permissions=True)


@frappe.whitelist()
def history(file_name: str, limit: int = 200):
	"""Riwayat edit satu file, dikelompokkan per Save (terbaru dulu)."""
	_get_file(file_name)
	rows = frappe.get_all(
		"CRM Spreadsheet Log",
		filters={"file": file_name},
		fields=["batch", "sheet", "cell", "old_value", "new_value", "owner", "creation"],
		order_by="creation desc",
		limit_page_length=cint(limit) or 200,
	)

	batches = {}
	for r in rows:
		b = batches.setdefault(
			r.batch,
			{
				"batch": r.batch,
				"sheet": r.sheet,
				"owner": r.owner,
				"owner_name": frappe.get_cached_value("User", r.owner, "full_name") or r.owner,
				"creation": r.creation,
				"cells": [],
			},
		)
		b["cells"].append({"cell": r.cell, "old": r.old_value, "new": r.new_value})

	return list(batches.values())

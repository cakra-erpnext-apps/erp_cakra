import io

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils.file_manager import save_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from crm_cakra.api import spreadsheet as sp


def _sample_xlsx():
	"""Template mirip rekap tender: judul merge + berwarna, header bergaris, rupiah, rumus."""
	wb = Workbook()
	ws = wb.active
	ws.title = "RAB"
	ws.merge_cells("A1:C1")
	ws["A1"] = "REKAP ANGGARAN TENDER"
	ws["A1"].font = Font(bold=True, size=14, color="FFFF0000")
	ws["A1"].alignment = Alignment(horizontal="center")
	ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")

	thin = Side(style="thin", color="FF000000")
	for i, h in enumerate(["Item", "Qty", "Harga"], start=1):
		c = ws.cell(row=2, column=i, value=h)
		c.font = Font(bold=True)
		c.border = Border(top=thin, bottom=thin, left=thin, right=thin)

	ws["A3"] = "Truk Trailer"
	ws["B3"] = 3
	ws["C3"] = 1500000
	ws["C3"].number_format = '"Rp"#,##0.00'
	ws["A4"] = "Total"
	ws["C4"] = "=B3*C3"
	ws.column_dimensions["A"].width = 30

	buf = io.BytesIO()
	wb.save(buf)
	return buf.getvalue()


class TestSpreadsheet(FrappeTestCase):
	def setUp(self):
		self.file = save_file("tender_test.xlsx", _sample_xlsx(), None, None, is_private=1)

	def tearDown(self):
		frappe.delete_doc("File", self.file.name, force=True, ignore_permissions=True)

	def test_read_keeps_excel_formatting(self):
		d = sp.read(self.file.name)
		self.assertEqual(d["sheets"], ["RAB"])
		self.assertIn({"r": 1, "c": 1, "rs": 1, "cs": 3}, d["merges"])
		self.assertEqual(d["cells"][0][1], 0)  # sel yang ditelan merge

		a1 = d["cells"][0][0]
		self.assertEqual(a1["v"], "REKAP ANGGARAN TENDER")
		self.assertEqual(a1["s"]["b"], 1)
		self.assertEqual(a1["s"]["c"], "#FF0000")
		self.assertEqual(a1["s"]["bg"], "#FFFF00")
		self.assertEqual(a1["s"]["ha"], "center")

		self.assertEqual(d["cells"][1][0]["s"]["bd"]["top"], "#000000")
		self.assertEqual(d["cells"][2][2]["v"], "Rp1,500,000.00")
		self.assertEqual(d["cells"][2][2]["s"]["ha"], "right")
		self.assertEqual(d["cells"][3][2]["f"], "=B3*C3")
		self.assertEqual(d["widths"][0], round(30 * 7 + 5))

	def test_save_writes_back_without_losing_file(self):
		d = sp.save(
			self.file.name,
			"RAB",
			[
				{"r": 3, "c": 2, "v": "5"},
				{"r": 5, "c": 1, "v": "Catatan"},
				{"r": 5, "c": 2, "v": "08123456789"},
			],
		)
		self.assertEqual(d["cells"][2][1]["v"], "5")
		self.assertEqual(d["cells"][4][0]["v"], "Catatan")
		self.assertEqual(d["cells"][4][1]["v"], "08123456789")  # awalan 0 tetap teks
		self.assertEqual(d["cells"][3][2]["f"], "=B3*C3")  # rumus selamat
		self.assertEqual(d["cells"][0][0]["s"]["b"], 1)  # style selamat

		ws = load_workbook(frappe.get_doc("File", self.file.name).get_full_path())["RAB"]
		self.assertEqual(ws["B3"].value, 5)
		self.assertIsInstance(ws["B5"].value, str)

	def test_rejects_non_xlsx(self):
		f = save_file("catatan.txt", b"halo", None, None, is_private=1)
		self.addCleanup(frappe.delete_doc, "File", f.name, force=True, ignore_permissions=True)
		self.assertRaises(frappe.ValidationError, sp.read, f.name)


class TestTender(FrappeTestCase):
	def setUp(self):
		self.tender = frappe.get_doc(
			{
				"doctype": "CRM Tender",
				"subject": "Pengadaan Jasa Angkutan",
				"status": "Draft",
				"assigned_to": "Administrator",
			}
		).insert()

	def tearDown(self):
		frappe.delete_doc("CRM Tender", self.tender.name, force=True, ignore_permissions=True)

	def test_attachments_put_excel_first(self):
		from crm_cakra.fcrm.doctype.crm_tender.crm_tender import get_attachments

		save_file("ttd.txt", b"dummy", "CRM Tender", self.tender.name, is_private=1)
		save_file("penawaran.xlsx", _sample_xlsx(), "CRM Tender", self.tender.name, is_private=1)

		att = get_attachments(self.tender.name)
		self.assertEqual([a["file_name"].split(".")[-1] for a in att], ["xlsx", "txt"])
		self.assertTrue(att[0]["editable"])
		self.assertFalse(att[1]["editable"])

	def test_history_records_cell_diffs(self):
		f = save_file("rab.xlsx", _sample_xlsx(), "CRM Tender", self.tender.name, is_private=1)

		sp.save(f.name, "RAB", [{"r": 3, "c": 2, "v": "5"}])
		sp.save(f.name, "RAB", [{"r": 3, "c": 2, "v": "5"}])  # tanpa perubahan -> tak dicatat
		sp.save(f.name, "RAB", [{"r": 3, "c": 2, "v": "8"}])

		h = sp.history(f.name)
		self.assertEqual(len(h), 2)
		self.assertEqual(h[0]["cells"], [{"cell": "B3", "old": "5", "new": "8"}])
		self.assertEqual(h[1]["cells"], [{"cell": "B3", "old": "3", "new": "5"}])
		self.assertEqual(h[0]["sheet"], "RAB")

	def test_log_does_not_block_delete(self):
		"""Riwayat memakai Data, bukan Link: hapus lampiran/tender tidak boleh tertahan."""
		f = save_file("rab.xlsx", _sample_xlsx(), "CRM Tender", self.tender.name, is_private=1)
		sp.save(f.name, "RAB", [{"r": 3, "c": 2, "v": "9"}])

		frappe.delete_doc("File", f.name, force=True, ignore_permissions=True)
		self.assertEqual(
			frappe.db.count("CRM Spreadsheet Log", {"file": f.name}),
			1,
			"riwayat harus tetap ada setelah lampirannya dihapus",
		)

	def test_activities_timeline_works(self):
		from crm_cakra.api.activities import get_activities

		versions, _calls, _notes, _tasks, attachments = get_activities(self.tender.name)
		self.assertIn("creation", {a["activity_type"] for a in versions})
		self.assertEqual(attachments, [])

	def test_data_fields_layout_grid(self):
		"""Form New Tender merender layout ini juga; kalau hilang, grid diam-diam
		jatuh ke get_default_layout (2 kolom) di kedua tempat."""
		from crm_cakra.fcrm.doctype.crm_fields_layout.crm_fields_layout import get_fields_layout

		tabs = get_fields_layout("CRM Tender", "Data Fields")
		grid = [
			[[f["fieldname"] for f in col["fields"]] for col in section["columns"]]
			for section in tabs[0]["sections"]
		]
		self.assertEqual(grid[0], [["subject"], ["tender_type"], ["organization"], ["assigned_to"]])
		self.assertEqual(
			grid[1], [["issue_date"], ["closing_date"], ["result_date"], ["contract_period"]]
		)
		self.assertEqual(grid[2], [["estimation_value"], ["currency"], ["status"]])
